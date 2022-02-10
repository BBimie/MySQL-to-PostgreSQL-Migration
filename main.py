from email.policy import default
import pandas as pd
import pymysql
from typing import Optional
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.engine import Connection as PostgresConnection
from pymysql.connections import Connection as MysqlConnection
from pymysql.err import OperationalError
from decouple import config

#environment variables
MYSQL_DB_NAME = config('MYSQL_DB_NAME', default="batest")
POSTGRESQL_DB_NAME = config('POSTGRESQL_DB_NAME', default="batest_postgresql")
password = config('password')
MYSQL_DB_USER=config('MYSQL_DB_USER', default="root")
POSTGRESQL_DB_USER=config('POSTGRESQL_DB_USER', default="postgres")
DB_HOST='localhost'


def to_camel_case(col_name: str) -> str:
    """function to convert column names to camel case"""
    if '_' in col_name:
        components = col_name.split('_')
    elif ' ' in col_name:
        components = col_name.split(' ')
    else:
        return col_name
    # We capitalize the first letter of each component except the first one
    # with the 'title' method and join them together.
    return components[0] + ''.join(x.title() for x in components[1:])


def get_mysql_connection(database_name: Optional[str]=None) -> MysqlConnection:
    """function to get mysql connection"""
    if database_name is None:
        database_name = ""
    return pymysql.connect(host=DB_HOST,port=int(3306),user=MYSQL_DB_USER,passwd=password,db=database_name)

def create_mysql_database(connection: MysqlConnection) -> None:
    #create mysql database
    connection.cursor().execute("DROP DATABASE IF EXISTS {};".format(MYSQL_DB_NAME))
    connection.cursor().execute("CREATE DATABASE {};".format(MYSQL_DB_NAME))
    connection.close()

def get_postgresql_connection(database_name: Optional[str] = None) -> PostgresConnection:
    """function to get postgresql connection"""
    if database_name is None:
        database_name = ""
    engine = create_engine(f'postgresql://{POSTGRESQL_DB_USER}:{password}@{DB_HOST}/{database_name}', echo=False, isolation_level="AUTOCOMMIT")
    return engine.connect()


def create_postgres_database(connection: PostgresConnection) -> None:
    """function to create postgresql database"""
    connection.execute("DROP DATABASE IF EXISTS {};".format(POSTGRESQL_DB_NAME))
    connection.execute("CREATE DATABASE {};".format(POSTGRESQL_DB_NAME)) 
    connection.close()

def migrate_data_from_mysql_to_postgres(mysql_conn: MysqlConnection, postgres_conn: PostgresConnection) -> None:
    """function to migrate data from mysql to postgresql"""
    tables_mysql = pd.read_sql_query("SHOW TABLES", mysql_conn)

    for table in tables_mysql['Tables_in_{}'.format(MYSQL_DB_NAME)]:
        table_chunks = pd.read_sql_query("SELECT * FROM {}".format(table), mysql_conn, chunksize=100)
        
        for chunk in table_chunks:
            table_cols = chunk.columns
            
            new_col_names = []
            for col_name in table_cols:
                new_col_names.append(to_camel_case(col_name))
            chunk.columns = new_col_names
            
            print(chunk.columns)
            print('====')
            
            #load chunk into postgresql database
            chunk.to_sql('{}'.format(table), con=postgres_conn, if_exists='append', index=False)


def write_to_excel(postgres_conn: PostgresConnection) -> None:
   tables_post = pd.read_sql_query("SELECT table_name FROM information_schema.tables WHERE table_schema='public'", postgres_conn)

   tables_post.to_excel('output.xlsx', sheet_name='{}'.format('All tables'), index=False)
   
   for table in tables_post['table_name']:
        data = pd.read_sql_query("SELECT * FROM public.\"{}\"".format(table), postgres_conn)
        
        book = load_workbook('output.xlsx')
        writer = pd.ExcelWriter('output.xlsx', engine='openpyxl')
        writer.book = book

        # copy existing sheets
        writer.sheets = dict((ws.title, ws) for ws in writer.book.worksheets)
        
        data.to_excel(writer, '{}'.format(table), index = False)

        writer.save()

def load_data_dump(filename: str, mysql_conn: MysqlConnection) -> None:
    fd = open(filename, 'r')
    sql_file = fd.read()
    fd.close()

    # all SQL commands (split on ';')
    queries = sql_file.split(';')

    # Execute every command from the input file
    for command in queries:
        #print(command)
        try:
            mysql_conn.cursor().execute(command)
        except OperationalError:
            pass

def main() -> None:
    mysql_conn = get_mysql_connection()
    
    #read sql script to create database
    #sql_file = open('batest.sql', 'rb')
    #sql_as_string = sql_file.read().decode('utf-8')

    #create mysql database
    create_mysql_database(mysql_conn)

    #Re-initialize mysql connection to create database
    mysql_conn = get_mysql_connection(MYSQL_DB_NAME)

    #load mysql database
    load_data_dump('batest.sql', mysql_conn)

    # INitialize postgres connection without database name and then create database
    postgres_conn = get_postgresql_connection()

    create_postgres_database(postgres_conn)

    # Re-initialize postgres connection with database name to perform operations on it
    postgres_conn = get_postgresql_connection(POSTGRESQL_DB_NAME)

    migrate_data_from_mysql_to_postgres(mysql_conn, postgres_conn)

    write_to_excel(postgres_conn)


if __name__ == '__main__':
    main()
